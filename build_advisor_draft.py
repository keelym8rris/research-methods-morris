"""Build a readable advisor DOCX from the audited Markdown and Python figures.
IEEE-style numbered references; single-column review layout, not a venue template.
"""
from pathlib import Path
import argparse,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
ROOT=Path(__file__).resolve().parent

def inline(p,text):
    # Preserve numbered references; create actual Word hyperlinks.
    pattern=r'\[([^\]]+)\]\((https?://[^)]+)\)'
    pos=0
    for m in re.finditer(pattern,text):
        p.add_run(text[pos:m.start()].replace('`',''))
        rel=p.part.relate_to(m.group(2),'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
        h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel)
        r=OxmlElement('w:r');pr=OxmlElement('w:rPr');col=OxmlElement('w:color');col.set(qn('w:val'),'174C70');pr.append(col);r.append(pr)
        t=OxmlElement('w:t');t.text=m.group(1);r.append(t);h.append(r);p._p.append(h);pos=m.end()
    p.add_run(text[pos:].replace('`',''))

def build(output):
    d=Document();s=d.sections[0]
    s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.top_margin=s.bottom_margin=Inches(.7);s.left_margin=s.right_margin=Inches(.8)
    for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
        st=d.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        fonts=st.element.get_or_add_rPr().get_or_add_rFonts()
        for key in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
            fonts.attrib.pop(qn('w:'+key),None)
        for key in ('ascii','hAnsi','eastAsia','cs'):fonts.set(qn('w:'+key),'Times New Roman')
    for border in d.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    n=d.styles['Normal'];n.font.size=Pt(11);n.paragraph_format.space_after=Pt(5);n.paragraph_format.line_spacing=1.05
    for name,size in [('Heading 1',12),('Heading 2',11)]:
        st=d.styles[name];st.font.size=Pt(size);st.font.bold=True;st.paragraph_format.space_before=Pt(9);st.paragraph_format.space_after=Pt(4);st.paragraph_format.keep_with_next=True
    d.styles['Title'].font.size=Pt(18);d.styles['Title'].font.bold=True
    d.styles['Caption'].font.size=Pt(9);d.styles['Caption'].font.italic=False
    # Page numbers support advisor comments on a longer draft.
    p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
    lines=(ROOT/'docs/thesis_draft.md').read_text().splitlines();i=0;refs=False
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line.startswith('# '):d.add_paragraph(line[2:],style='Title');continue
        if line.startswith('## '):
            title=line[3:];refs=title=='References'
            p=d.add_paragraph(title,style='Heading 1')
            if refs:p.paragraph_format.page_break_before=True
            continue
        if line.startswith('### '):d.add_paragraph(line[4:],style='Heading 2');continue
        if line.startswith('!['):
            m=re.match(r'!\[([^\]]+)\]\(([^)]+)\)',line)
            p=d.add_paragraph();p.paragraph_format.keep_with_next=True
            pic=p.add_run().add_picture(str((ROOT/'docs'/m.group(2)).resolve()),width=Inches(6.7))
            pic._inline.docPr.set('descr',m.group(1)+' generated from audited held-out predictions and metrics')
            continue
        if line.startswith('|'):
            rows=[[c.strip() for c in line.strip('|').split('|')]]
            while i<len(lines) and lines[i].strip().startswith('|'):
                l=lines[i].strip();i+=1
                if re.match(r'^\|[-:| ]+\|$',l):continue
                rows.append([c.strip() for c in l.strip('|').split('|')])
            table=d.add_table(rows=0,cols=len(rows[0]));table.style='Table Grid'
            widths={2:[2.0,4.9],3:[1.2,2.8,2.9],4:[2.1,1.6,1.6,1.6]}[len(rows[0])]
            table.autofit=False
            for j,width in enumerate(widths):table.columns[j].width=Inches(width)
            for k,vals in enumerate(rows):
                cells=table.add_row().cells
                for j,v in enumerate(vals):
                    cells[j].width=Inches(widths[j]);p=cells[j].paragraphs[0];p.paragraph_format.space_after=Pt(3)
                    inline(p,v)
                    for r in p.runs:r.font.size=Pt(9);r.bold=k==0
                trpr=table.rows[-1]._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');trpr.append(no)
                if k==0:trpr.append(OxmlElement('w:tblHeader'))
            d.add_paragraph().paragraph_format.space_after=Pt(1);continue
        p=d.add_paragraph(style='Caption' if line.startswith(('Figure ','Table I.')) else 'Normal')
        inline(p,line)
        if refs:
            p.paragraph_format.left_indent=Inches(.25);p.paragraph_format.first_line_indent=Inches(-.25)
            for r in p.runs:r.font.size=Pt(10)
    d.core_properties.title=lines[0][2:];d.core_properties.author='Keely Morris'
    d.core_properties.subject='Prostate data analysis and proposed screening research'
    output=Path(output);output.parent.mkdir(parents=True,exist_ok=True);d.save(output);print(output)
if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--output',type=Path,default=ROOT.parent/'deliverables/Keely_Morris_IEEE_Thesis_Draft.docx');build(p.parse_args().output)
