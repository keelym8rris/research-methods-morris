"""Build the thesis working paper as an IEEE-style two-column DOCX."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "docs" / "thesis_draft.md"
OUTPUT = ROOT.parent / "Keely_Morris_IEEE_Thesis_Draft.docx"

FONT = "Times New Roman"


def set_run_font(run, size=None, bold=None, italic=None, small_caps=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if small_caps is not None:
        run.font.small_caps = small_caps


def add_inline(paragraph, text, size=10):
    """Add a small subset of Markdown inline formatting."""
    pattern = re.compile(r"(\*\*.*?\*\*|\*.*?\*|`.*?`)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            set_run_font(paragraph.add_run(text[cursor:match.start()]), size=size)
        token = match.group(0)
        if token.startswith("**"):
            set_run_font(paragraph.add_run(token[2:-2]), size=size, bold=True)
        elif token.startswith("*"):
            set_run_font(paragraph.add_run(token[1:-1]), size=size, italic=True)
        else:
            set_run_font(paragraph.add_run(token[1:-1]), size=size)
        cursor = match.end()
    if cursor < len(text):
        set_run_font(paragraph.add_run(text[cursor:]), size=size)


def set_cell_margins(cell, top=50, start=65, bottom=50, end=65):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(0.625)
    section.right_margin = Inches(0.625)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.4)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(10)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.first_line_indent = Inches(0.15)
    normal.paragraph_format.widow_control = True

    h1 = doc.styles["Heading 1"]
    h1.font.name = FONT
    h1._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    h1._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    h1.font.size = Pt(10)
    h1.font.bold = False
    h1.font.small_caps = True
    h1.font.color.rgb = RGBColor(0, 0, 0)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.space_before = Pt(8)
    h1.paragraph_format.space_after = Pt(3)
    h1.paragraph_format.keep_with_next = True
    h1.paragraph_format.first_line_indent = Inches(0)

    h2 = doc.styles["Heading 2"]
    h2.font.name = FONT
    h2._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    h2._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    h2.font.size = Pt(10)
    h2.font.bold = False
    h2.font.italic = True
    h2.font.color.rgb = RGBColor(0, 0, 0)
    h2.paragraph_format.space_before = Pt(6)
    h2.paragraph_format.space_after = Pt(2)
    h2.paragraph_format.keep_with_next = True
    h2.paragraph_format.first_line_indent = Inches(0)

    ref = doc.styles.add_style("IEEE Reference", 1)
    ref.font.name = FONT
    ref._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    ref._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    ref.font.size = Pt(8)
    ref.paragraph_format.line_spacing = 1.0
    ref.paragraph_format.space_after = Pt(1.5)
    ref.paragraph_format.left_indent = Inches(0.17)
    ref.paragraph_format.first_line_indent = Inches(-0.17)
    ref.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    ref.paragraph_format.widow_control = True


def set_two_columns(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(0.625)
    section.right_margin = Inches(0.625)
    sect_pr = section._sectPr
    cols = sect_pr.xpath("./w:cols")
    if cols:
        cols = cols[0]
    else:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), "2")
    cols.set(qn("w:space"), "360")
    cols.set(qn("w:equalWidth"), "1")


def add_title_block(doc, lines):
    title = lines[0][2:].strip()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    set_run_font(p.add_run(title), size=24)

    author = doc.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.paragraph_format.space_after = Pt(1)
    set_run_font(author.add_run(lines[2].strip()), size=11)

    for text in [lines[3].strip(), lines[4].strip(), lines[5].strip()]:
        text = text.rstrip("  ")
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(text), size=9.5, italic=text.startswith("Revised"))


def add_abstract(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_run_font(p.add_run("Abstract—"), size=9, bold=True, italic=True)
    set_run_font(p.add_run(text), size=9, italic=True)


def add_index_terms(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.space_after = Pt(6)
    set_run_font(p.add_run("Index Terms—"), size=9, bold=True, italic=True)
    cleaned = re.sub(r"^\*\*Index Terms—\*\*", "", text).strip()
    set_run_font(p.add_run(cleaned), size=9, italic=True)


def add_results_table(doc, rows):
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Inches(0)
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(2)
    set_run_font(caption.add_run("TABLE I\nREPEATED CROSS-VALIDATION PERFORMANCE (50 EVALUATIONS PER MODEL)"), size=8, small_caps=True)

    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r_idx, values in enumerate(rows):
        for c_idx, value in enumerate(values):
            cell = table.cell(r_idx, c_idx)
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            set_run_font(p.add_run(value), size=7.2, bold=r_idx == 0)
    header_pr = table.rows[0]._tr.get_or_add_trPr()
    header_flag = OxmlElement("w:tblHeader")
    header_flag.set(qn("w:val"), "true")
    header_pr.append(header_flag)
    set_table_geometry(table, [1740, 1100, 1100, 1100])
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(0)
    after.paragraph_format.first_line_indent = Inches(0)


def parse_source(doc):
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    add_title_block(doc, lines)

    index = 6
    while index < len(lines) and not lines[index].startswith("## Abstract"):
        index += 1
    index += 1
    while index < len(lines) and not lines[index].strip():
        index += 1
    abstract_text = lines[index].strip()
    add_abstract(doc, abstract_text)
    index += 1
    while index < len(lines) and not lines[index].strip():
        index += 1
    add_index_terms(doc, lines[index].strip())
    index += 1

    body_section = doc.add_section(WD_SECTION.CONTINUOUS)
    set_two_columns(body_section)
    in_references = False

    while index < len(lines):
        raw = lines[index]
        line = raw.strip()
        if not line:
            index += 1
            continue
        if line.startswith("## "):
            heading = line[3:].strip()
            in_references = heading == "References"
            doc.add_paragraph(heading.upper() if not in_references else "REFERENCES", style="Heading 1")
            index += 1
            continue
        if line.startswith("### "):
            doc.add_paragraph(line[4:].strip(), style="Heading 2")
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and re.match(r"^\|[-:| ]+\|$", lines[index + 1].strip()):
            table_rows = []
            table_rows.append([x.strip() for x in line.strip("|").split("|")])
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_rows.append([x.strip() for x in lines[index].strip().strip("|").split("|")])
                index += 1
            add_results_table(doc, table_rows)
            continue

        style = "IEEE Reference" if in_references and re.match(r"^\[\d+\]", line) else "Normal"
        p = doc.add_paragraph(style=style)
        if style == "Normal":
            p.paragraph_format.first_line_indent = Inches(0.15)
        add_inline(p, line, size=8 if style == "IEEE Reference" else 10)
        index += 1


def build():
    doc = Document()
    configure_document(doc)
    parse_source(doc)
    doc.core_properties.title = "Responsible and Interpretable Machine Learning for PSA Prediction and Prostate Cancer Risk Modeling"
    doc.core_properties.author = "Keely Morris"
    doc.core_properties.subject = "IEEE-style undergraduate research working paper"
    doc.core_properties.comments = "Revised July 2026; 24 references; two-column IEEE-style layout."
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
